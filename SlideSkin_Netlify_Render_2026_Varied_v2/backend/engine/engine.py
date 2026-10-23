"""SlideSkin beta v2: diversified template-native layouts and mandatory bookends.
Experimental: native charts/tables/media are flagged rather than silently certified.
"""
import io, json, math, re
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from base_engine import (source_content, clone_template_slide, set_text,
                         get_title, get_bodies, EMU)

COVER=8
CLOSING=28
TEXT_LAYOUTS=[9,10,11,14,15,21,22,23]
IMAGE_LAYOUTS=[12,16,17,18]

def is_intro(c):
    t=c['title'].lower().strip()
    return bool(re.search(r'^(welcome|introduction|overview|title slide|opening)(\b|$)',t))

def is_closing(c):
    t=c['title'].lower().strip()
    return bool(re.search(r'^(thank you|thanks|questions|q\s*&\s*a|the end|closing)(\b|$)',t))

def choose_layout(c,previous,usage,slide_num):
    imgs=len(c['images']); n=len(c['bullets']); chars=sum(map(len,c['bullets']))
    candidates=IMAGE_LAYOUTS if imgs else TEXT_LAYOUTS
    scores=[]
    for idx in candidates:
        score=0.0
        if imgs:
            score += {12:4,16:3,17:5,18:5}[idx]
            if n>=4: score += {12:1,16:1,17:3,18:3}[idx]
            if chars>500:score += {12:2,16:2,17:-3,18:-3}[idx]
        else:
            if n>=5:score += {9:3,10:4,11:3,14:4,15:3,21:0,22:0,23:0}[idx]
            elif n>=3:score += {9:3,10:3,11:3,14:3,15:4,21:4,22:4,23:3}[idx]
            else:score += {9:4,10:2,11:3,14:4,15:2,21:3,22:3,23:3}[idx]
            if chars>650:score += 2 if idx in (9,10,14) else -3
        score-=usage.get(idx,0)*1.25
        if idx==previous:score-=7
        if idx in (21,22) and previous in (21,22):score-=3
        # Deterministic tie break with mild rotation.
        score += ((idx*7+slide_num*3)%11)*.015
        scores.append((score,idx))
    return max(scores)[1]

def notes_copy(src,dst):
    try:
        value=src.notes_slide.notes_text_frame.text
        if value:dst.notes_slide.notes_text_frame.text=value
    except Exception:pass

def put_image_contained(slide,src_shape,box):
    from PIL import Image
    data=src_shape.image.blob
    with Image.open(io.BytesIO(data)) as img: iw,ih=img.size
    x,y,w,h=box
    ratio=min(w/iw,h/ih)
    tw=int(iw*ratio);th=int(ih*ratio)
    slide.shapes.add_picture(io.BytesIO(data),x+(w-tw)//2,y+(h-th)//2,width=tw,height=th)

def add_cover(deck,originals,title,subtitle,idx,src=None):
    dest=clone_template_slide(deck,originals[idx-1])
    titles=[s for s in dest.shapes if s.has_text_frame and 'title' in s.name.lower() and 'subtitle' not in s.name.lower()]
    subs=[s for s in dest.shapes if s.has_text_frame and 'subtitle' in s.name.lower()]
    if titles:set_text(titles[0],title)
    if subs:set_text(subs[0],subtitle)
    if src:notes_copy(src,dest)
    return dest

def populate(slide,c,src,page,idx):
    warnings=[]
    title=get_title(slide)
    if title:set_text(title,c['title'])
    for s in slide.shapes:
        if s.has_text_frame and 'slide number' in s.name.lower():set_text(s,str(page))
    bodies=get_bodies(slide)
    # Text placeholders are filled in-place; sample content must not survive.
    for s in bodies:set_text(s,'')
    payload='\n'.join(([c['subtitle']] if c['subtitle'] else [])+['• '+b for b in c['bullets']])
    if bodies:
        # The largest usable body is preferred, with smaller regions cleared.
        target=max(bodies,key=lambda s:s.width*s.height)
        if len(payload)>950:warnings.append('Dense content: inspect text overflow visually')
        size=17 if len(payload)>650 else (18 if len(payload)>400 else 19)
        set_text(target,payload,size)
    elif payload:
        from base_engine import add_text
        add_text(slide,1.1,2.0,10.7,3.9,payload,18)
        warnings.append('No native text placeholder; used template-derived textbox')
    # Use template picture placeholder geometry rather than arbitrary image positions.
    image_boxes=[]
    for sh in list(slide.shapes):
        if sh is title:continue
        if ('picture placeholder' in sh.name.lower() or
            (sh.is_placeholder and sh.placeholder_format.type==18)):
            image_boxes.append((sh.left,sh.top,sh.width,sh.height))
            sh._element.getparent().remove(sh._element)
    if c['images']:
        if not image_boxes:
            # For slides where picture regions exist in the layout but not slide XML.
            image_boxes=[(int(7.45*EMU),int(1.85*EMU),int(4.65*EMU),int(4.55*EMU))]
            if bodies:
                target=max(bodies,key=lambda s:s.width*s.height)
                right=image_boxes[0][0]-int(.25*EMU)
                if target.left+target.width>right:
                    target.width=max(int(2.5*EMU),right-target.left)
            warnings.append('Picture region inferred; inspect placement')
        if len(c['images'])>len(image_boxes):
            warnings.append('Additional images arranged in primary image region')
        x,y,w,h=image_boxes[0]
        count=len(c['images'])
        cols=math.ceil(math.sqrt(count));rows=math.ceil(count/cols)
        gap=int(.09*EMU)
        for j,img in enumerate(c['images']):
            col=j%cols;row=j//cols
            cellw=(w-gap*(cols-1))//cols;cellh=(h-gap*(rows-1))//rows
            put_image_contained(slide,img,(x+col*(cellw+gap),y+row*(cellh+gap),cellw,cellh))
    for key in ('charts','tables','media'):
        if c[key]:warnings.append(f'{len(c[key])} native {key} object(s) not yet supported by fidelity-preserving export')
    notes_copy(src,slide)
    return warnings

def run(template,source,output,report_path=None,min_pt=17):
    deck=Presentation(str(template));incoming=Presentation(str(source));originals=list(deck.slides)
    if len(originals)<29:raise ValueError('Expected 2026 All Hands template with 29 slides')
    contents=[source_content(s) for s in incoming.slides]
    report=[];usage={};previous=None;page=0
    # Guaranteed branded introduction, unless source begins with a clear introduction.
    intro_source=0 if contents and is_intro(contents[0]) else None
    if intro_source is None:
        title=contents[0]['title'] if contents else 'Presentation'
        add_cover(deck,originals,title,'Presented with SlideSkin',COVER)
        page+=1
        report.append({'source_slide':None,'output_slides':[page],'template_slide':COVER,'mode':'generated_introduction','warnings':[]})
    for i,(src,c) in enumerate(zip(incoming.slides,contents)):
        if i==0 and is_intro(c):
            page+=1;add_cover(deck,originals,c['title'],c['subtitle'],COVER,src)
            report.append({'source_slide':i+1,'output_slides':[page],'template_slide':COVER,'mode':'source_introduction','warnings':[]})
            continue
        if i==len(contents)-1 and is_closing(c):
            # Do not discard any content from a closing slide: it can have images/body.
            if c['images'] or c['bullets']:
                idx=choose_layout(c,previous,usage,i)
                page+=1;dst=clone_template_slide(deck,originals[idx-1]);warnings=populate(dst,c,src,page,idx)
                report.append({'source_slide':i+1,'output_slides':[page],'template_slide':idx,'mode':'source_closing_content','warnings':warnings})
            else:
                page+=1;add_cover(deck,originals,c['title'],c['subtitle'],CLOSING,src)
                report.append({'source_slide':i+1,'output_slides':[page],'template_slide':CLOSING,'mode':'source_closing','warnings':[]})
            continue
        idx=choose_layout(c,previous,usage,i)
        usage[idx]=usage.get(idx,0)+1;previous=idx
        page+=1;dst=clone_template_slide(deck,originals[idx-1]);warnings=populate(dst,c,src,page,idx)
        report.append({'source_slide':i+1,'output_slides':[page],'template_slide':idx,'mode':'content_matched','warnings':warnings})
    # Always finish with a branded closing, unless a simple native closing already exists.
    if not(contents and is_closing(contents[-1]) and not contents[-1]['images'] and not contents[-1]['bullets']):
        page+=1;add_cover(deck,originals,'Thank You!','Questions?',CLOSING)
        report.append({'source_slide':None,'output_slides':[page],'template_slide':CLOSING,'mode':'generated_closing','warnings':[]})
    for _ in originals:
        sid=deck.slides._sldIdLst[0];deck.part.drop_rel(sid.rId);deck.slides._sldIdLst.remove(sid)
    deck.save(str(output))
    if report_path:Path(report_path).write_text(json.dumps(report,indent=2))
    return report
