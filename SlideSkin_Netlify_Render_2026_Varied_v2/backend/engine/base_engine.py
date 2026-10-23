"""SlideSkin v2: adaptive, template-derived editable PowerPoint layout engine."""
import argparse, copy, json, math, re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import _Relationship

EMU=914400

def text_shapes(slide):
    return [x for x in slide.shapes if x.has_text_frame and x.text.strip()]

def source_content(slide):
    shapes=text_shapes(slide)
    texts=[x.text.strip() for x in shapes]
    title=slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.text.strip() else (texts[0] if texts else '')
    remaining=[v for v in texts if not re.fullmatch(r'\d{1,3}\s*/\s*\d{1,3}',v)]
    if title in remaining: remaining.remove(title)
    subtitle=remaining.pop(0) if remaining else ''
    bullets=[]
    for t in remaining:
        bullets.extend([v.strip().lstrip('•').strip() for v in t.splitlines() if v.strip()])
    # Some sources store bullets in separate shapes; do not lose any.
    imgs=[s for s in slide.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE]
    charts=[s for s in slide.shapes if s.has_chart]
    tables=[s for s in slide.shapes if s.has_table]
    media=[s for s in slide.shapes if 'media' in s.name.lower() or s.shape_type==MSO_SHAPE_TYPE.MEDIA]
    return dict(title=title,subtitle=subtitle,bullets=bullets,images=imgs,charts=charts,tables=tables,media=media,raw_text=texts)

def is_number(s): return bool(re.fullmatch(r'\d+',s.strip()))
def classify_template(slide,idx):
    ts=text_shapes(slide)
    titles=[s for s in ts if 'title' in s.name.lower()]
    content=[s for s in ts if ('content' in s.name.lower() or 'text placeholder' in s.name.lower()) and not is_number(s.text)]
    imgs=[s for s in slide.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE]
    charts=[s for s in slide.shapes if s.has_chart]
    return dict(index=idx,content_count=len(content),has_image=bool(imgs),has_chart=bool(charts),title=bool(titles),area=sum(s.width*s.height for s in content)/(EMU**2),content_shapes=content)

def select_layout(c, catalog, slide_index, total):
    if slide_index==0: return 3
    if slide_index==total-1 and ('thank' in c['title'].lower() or 'end' in c['title'].lower()): return 26
    n=len(c['bullets']); imgs=len(c['images']); charts=len(c['charts'])
    preferred=[]
    for v in catalog:
        i=v['index']
        if i in (1,2,3,4,5,26): continue
        score=0
        if charts: score += 8 if v['has_chart'] else -9
        elif v['has_chart']: score-=8
        if imgs: score += 6 if v['has_image'] else -5
        elif v['has_image']: score-=5
        if n in (3,4) and v['content_count']>=3: score+=8
        if n>4 and v['content_count']==1: score+=3
        if not n and v['content_count']==1: score+=3
        score+=min(v['area'],20)*.12
        if i in (7,8,10,11,14,17,21): score+=2
        if i in (15,16,19,25): score-=4
        preferred.append((score,i))
    return max(preferred)[1]

def clone_template_slide(deck, source_slide):
    dest=deck.slides.add_slide(source_slide.slide_layout)
    for sh in list(dest.shapes):
        el=sh._element; el.getparent().remove(el)
    for sh in source_slide.shapes:
        dest.shapes._spTree.insert_element_before(copy.deepcopy(sh._element),'p:extLst')
    for rel in source_slide.part.rels.values():
        if rel.reltype in (RT.SLIDE_LAYOUT,RT.NOTES_SLIDE): continue
        if rel.rId in dest.part.rels._rels: continue
        dest.part.rels._rels[rel.rId]=_Relationship(dest.part.rels._base_uri,rel.rId,rel.reltype,'External' if rel.is_external else 'Internal',rel.target_ref if rel.is_external else rel.target_part)
    return dest

def set_text(shape, value, font_size=None):
    tf=shape.text_frame
    # Retain placeholder formatting by editing the first existing run where possible.
    old=next((r for p in tf.paragraphs for r in p.runs),None)
    props=None
    if old is not None:
        props=(old.font.name,old.font.size,old.font.bold)
    tf.clear(); tf.word_wrap=True
    for j,line in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
        run=p.add_run(); run.text=line
        if props:
            if props[0]: run.font.name=props[0]
            if props[1]: run.font.size=props[1]
            if props[2] is not None: run.font.bold=props[2]
        if font_size: run.font.size=Pt(font_size)
    tf.auto_size=MSO_AUTO_SIZE.NONE

def approximate_capacity(shape, pt=19):
    w=shape.width/EMU; h=shape.height/EMU
    chars_per_line=max(10,int(w*72/(pt*.51)))
    lines=max(1,int(h*72/(pt*1.4)))
    return chars_per_line*lines

def add_text(slide,x,y,w,h,text,size=19,color=None):
    s=slide.shapes.add_textbox(int(x*EMU),int(y*EMU),int(w*EMU),int(h*EMU))
    set_text(s,text,size)
    if color:
        for p in s.text_frame.paragraphs:
            for r in p.runs:r.font.color.rgb=RGBColor(*color)
    return s

def get_title(slide):
    return next((s for s in slide.shapes if s.has_text_frame and 'title' in s.name.lower()),None)

def get_bodies(slide):
    return [s for s in slide.shapes if s.has_text_frame and ('content placeholder' in s.name.lower() or 'text placeholder' in s.name.lower()) and not is_number(s.text)]

def fill_slide(slide,c,page,source_slide,base_index):
    title=get_title(slide)
    if title:set_text(title,c['title'])
    for s in slide.shapes:
        if s.has_text_frame and ('slide number' in s.name.lower()):set_text(s,str(page))
    bodies=get_bodies(slide)
    # Clear ALL template example text from non-title content regions.
    for s in bodies:set_text(s,'')
    items=[c['subtitle']]+c['bullets'] if c['subtitle'] else c['bullets'][:]
    if base_index in (3,28):
        items=[]
        subs=[s for s in slide.shapes if s.has_text_frame and 'subtitle' in s.name.lower()]
        for s in (subs or [x for x in slide.shapes if x.has_text_frame and x is not title and x.text.strip()][:1]):set_text(s,c['subtitle'])
    elif len(bodies)>=6 and len(c['bullets'])>=3:
        # Three-column template: alternating headings and content; use subtitle once.
        ordered=sorted(bodies,key=lambda x:(x.left,x.top))
        cols=[]
        for s in ordered:
            near=next((group for group in cols if abs(group[0].left-s.left)<EMU*.5),None)
            if near is None:cols.append([s])
            else:near.append(s)
        for j,group in enumerate(cols[:3]):
            group.sort(key=lambda x:x.top)
            if group:set_text(group[0],c['subtitle'] if j==0 else f'POINT {j+1}')
            if len(group)>1:set_text(group[1],c['bullets'][j])
        items=c['bullets'][3:]
    elif bodies:
        # Adaptive text: fill primary existing content region.
        primary=max(bodies,key=lambda x:x.width*x.height)
        payload='\n'.join(([c['subtitle']] if c['subtitle'] else [])+['• '+b for b in c['bullets']])
        cap=approximate_capacity(primary,19)
        if len(payload)>cap:
            # Expand into additional template-derived columns, without shrinking below 17pt.
            left=primary.left/EMU; top=primary.top/EMU; width=primary.width/EMU; height=primary.height/EMU
            chunks=[]; current=''
            for line in payload.splitlines():
                if len(current)+len(line)>max(80,cap//2) and current:
                    chunks.append(current);current=''
                current+=(('\n' if current else '')+line)
            if current:chunks.append(current)
            if len(chunks)<=2 and width>7:
                primary.width=int((width-.35)/2*EMU)
                set_text(primary,chunks[0],18)
                add_text(slide,left+width/2+.18,top,(width-.35)/2,height,chunks[1] if len(chunks)>1 else '',18)
            else:
                set_text(primary,payload,17)
        else:set_text(primary,payload,19)
        items=[]
    # Handle remaining content without silent loss.
    if items and base_index not in (3,28):
        add_text(slide,.95,6.2,10.9,.5,' | '.join(items),12)
    # Source images: add into a template-derived content region; preserve image pixels.
    if c['images']:
        usable=[s for s in bodies if s.width/EMU>3]
        if usable:
            region=max(usable,key=lambda x:x.width*x.height)
            # Reserve right side for image gallery, reflow text to left.
            if region.width/EMU>7:
                left=region.left/EMU; top=region.top/EMU; w=region.width/EMU; h=region.height/EMU
                region.width=int(w*.52*EMU)
                iw=w*.43; ix=left+w*.57
                for j,im in enumerate(c['images'][:4]):
                    rows=math.ceil(min(len(c['images']),4)/2); cols=min(len(c['images']),2)
                    cellw=iw/cols; cellh=h/rows
                    slide.shapes.add_picture(__import__('io').BytesIO(im.image.blob),int((ix+j%cols*cellw)*EMU),int((top+j//cols*cellh)*EMU),width=int((cellw-.08)*EMU),height=int((cellh-.08)*EMU))
    # Charts and tables: editable native copies need relationship rewrites; mark for manual review instead of silently dropping.
    warnings=[]
    for k in ('charts','tables','media'):
        if c[k]:warnings.append(f'{len(c[k])} {k} object(s) require native-object preservation pass')
    try:
        notes=source_slide.notes_slide.notes_text_frame.text
        if notes and slide.notes_slide.notes_text_frame is not None:slide.notes_slide.notes_text_frame.text=notes
    except Exception:pass
    return warnings

def redesign(template,source,output,report_path=None):
    t=Presentation(str(template)); s=Presentation(str(source))
    originals=list(t.slides)
    catalog=[classify_template(sl,i+1) for i,sl in enumerate(originals)]
    report=[]
    for i,src in enumerate(s.slides):
        c=source_content(src)
        candidate=select_layout(c,catalog,i,len(s.slides))
        dest=clone_template_slide(t,originals[candidate-1])
        warnings=fill_slide(dest,c,i+1,src,candidate)
        report.append(dict(source_slide=i+1,template_slide=candidate,title=c['title'],bullet_count=len(c['bullets']),image_count=len(c['images']),warnings=warnings))
    for i in range(len(originals)-1,-1,-1):
        sid=t.slides._sldIdLst[0];t.part.drop_rel(sid.rId);t.slides._sldIdLst.remove(sid)
    t.save(str(output))
    if report_path:Path(report_path).write_text(json.dumps(report,indent=2))
    return report

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--template',required=True);ap.add_argument('--source',required=True);ap.add_argument('--output',required=True);ap.add_argument('--report');a=ap.parse_args()
    print(json.dumps(redesign(a.template,a.source,a.output,a.report),indent=2))
