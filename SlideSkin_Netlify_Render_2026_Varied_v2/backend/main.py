"""Experimental public API; in-memory file delivery, no persistent uploads."""
import io, json, os, sys, tempfile, threading, time, uuid
from pathlib import Path
from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from pptx import Presentation
sys.path.insert(0, str(Path(__file__).parent/'engine'))
from engine import run
ROOT=Path(__file__).parent
TEMPLATE=ROOT/'templates'/'all_hands_2026.pptx'
MAX_MB=int(os.environ.get('MAX_UPLOAD_MB','20'))
TTL=900
results={}
lock=threading.Lock()
app=FastAPI(title="Ray's Magic PowerPoint Converter API",version='0.1-beta')
origins=[v.strip().rstrip('/') for v in os.environ.get('FRONTEND_ORIGINS','http://localhost:8888,http://127.0.0.1:8888').split(',') if v.strip()]
app.add_middleware(CORSMiddleware,allow_origins=origins,allow_credentials=False,allow_methods=['GET','POST'],allow_headers=['*'])
@app.get('/api/health')
def health():return {'status':'ok','template':'All Hands 2026','mode':'rule-based beta'}
@app.post('/api/transform')
async def transform(source:UploadFile=File(...)):
    if not source.filename or not source.filename.lower().endswith('.pptx'):raise HTTPException(400,'Upload a .pptx source file.')
    data=await source.read(MAX_MB*1024*1024+1)
    if len(data)>MAX_MB*1024*1024:raise HTTPException(413,f'Maximum source size is {MAX_MB} MB.')
    if not data.startswith(b'PK'):raise HTTPException(400,'Invalid PowerPoint file.')
    if not lock.acquire(blocking=False):raise HTTPException(429,'Another presentation is processing. Please retry shortly.')
    try:
        with tempfile.TemporaryDirectory() as d:
            d=Path(d);sp=d/'source.pptx';op=d/'result.pptx';rp=d/'report.json';sp.write_bytes(data)
            try:
                nsrc=len(Presentation(sp).slides)
                if nsrc>50:raise HTTPException(400,'Beta limit: 50 slides per presentation.')
                mapping=run(TEMPLATE,sp,op,rp)
                output=op.read_bytes()
                nout=len(Presentation(io.BytesIO(output)).slides)
            except HTTPException:raise
            except Exception as exc:raise HTTPException(422,f'Could not transform this presentation: {str(exc)[:250]}')
        now=time.time()
        for k,v in list(results.items()):
            if v['created']+TTL<now:results.pop(k,None)
        key=uuid.uuid4().hex
        results[key]={'created':now,'pptx':output,'report':json.dumps(mapping,indent=2).encode()}
        return {'job_id':key,'source_slides':nsrc,'output_slides':nout,'mapping':mapping,'expires_in_seconds':TTL}
    finally:lock.release()
@app.get('/api/download/{job_id}/{kind}')
def download(job_id:str,kind:str):
    if kind not in ('pptx','report'):raise HTTPException(404,'Not found')
    result=results.get(job_id)
    if not result or result['created']+TTL<time.time():raise HTTPException(404,'Download expired. Please run again.')
    data=result['pptx' if kind=='pptx' else 'report']
    filename='Rays_Magic_PowerPoint_Converted.pptx' if kind=='pptx' else 'Rays_Magic_Conversion_Report.json'
    return Response(data,media_type='application/vnd.openxmlformats-officedocument.presentationml.presentation' if kind=='pptx' else 'application/json',headers={'Content-Disposition':f'attachment; filename="{filename}"','Cache-Control':'no-store'})
